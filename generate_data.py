import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import json
import os

# Lista de invitados para la boda en el Castillo de Santiago (Sanlúcar de Barrameda)
# Novios: Ester Muñoz & Samuel Madrid
invitados_data = [
    # Mesa Presidencial (0 o Presidencial) - Novios y padres
    (1, "Ester Muñoz", 0),
    (2, "Samuel Madrid", 0),
    (3, "Manuel Muñoz Salgado", 0),
    (4, "Carmen Barba Morales", 0),
    (5, "Francisco Madrid Peralta", 0),
    (6, "Isabel Gómez Ceballos", 0),

    # Mesa 1: Albariza (Familia Muñoz - Tíos y primos directos)
    (7, "Alfonso Muñoz Salgado", 1),
    (8, "María del Mar Prieto", 1),
    (9, "Gonzalo Muñoz Prieto", 1),
    (10, "Lucía Muñoz Prieto", 1),
    (11, "Javier Barba Morales", 1),
    (12, "Elena Moreno Rueda", 1),
    (13, "Pablo Barba Moreno", 1),
    (14, "Clara Barba Moreno", 1),

    # Mesa 2: Manzanilla (Familia Madrid - Tíos y primos paternos)
    (15, "Antonio Madrid Peralta", 2),
    (16, "Rocío Valero Domínguez", 2),
    (17, "Álvaro Madrid Valero", 2),
    (18, "Marta Madrid Valero", 2),
    (19, "Ignacio Madrid Peralta", 2),
    (20, "Teresa Cordero Gil", 2),
    (21, "Diego Madrid Cordero", 2),
    (22, "Beatriz Madrid Cordero", 2),

    # Mesa 3: Bajo de Guía (Familia Gómez - Tíos y primos maternos)
    (23, "Fernando Gómez Ceballos", 3),
    (24, "Esperanza Ramos Ortiz", 3),
    (25, "Carlos Gómez Ramos", 3),
    (26, "Laura Gómez Ramos", 3),
    (27, "Santiago Gómez Ceballos", 3),
    (28, "Pilar Delgado Blanco", 3),
    (29, "Manuel Gómez Delgado", 3),
    (30, "Sofía Gómez Delgado", 3),

    # Mesa 4: Doñana (Amigos íntimos Novia - Universidad)
    (31, "Valentina Castro Vega", 4),
    (32, "Mateo Herrera Ibáñez", 4),
    (33, "Daniela Morales Ponce", 4),
    (34, "Lucas Navas Serrano", 4),
    (35, "Camila Benítez Soler", 4),
    (36, "Adrián Fuentes Rivas", 4),
    (37, "Martina Vidal Sanz", 4),
    (38, "Nicolás Parra Méndez", 4),
    (39, "Jimena Reyes Gallardo", 4),

    # Mesa 5: Torre del Homenaje (Amigos íntimos Novio - Colegio y rugby)
    (40, "Jaime Lara Carrasco", 5),
    (41, "Guillermo Ortiz Marín", 5),
    (42, "Rodrigo Silva Campos", 5),
    (43, "Marcos Lozano Peña", 5),
    (44, "Bruno Quintana Esteban", 5),
    (45, "Inés Cobo Hidalgo", 5),
    (46, "Carlota Pardo Varela", 5),
    (47, "Tomás Cabrera Rojo", 5),
    (48, "Andrés Pastor Crespo", 5),

    # Mesa 6: Guadalquivir (Compañeros de trabajo Novia)
    (49, "Héctor Santos Miranda", 6),
    (50, "Natalia Guerra Molina", 6),
    (51, "Rubén Aguilar Cruz", 6),
    (52, "Silvia Cano Ferrer", 6),
    (53, "Víctor Montero Roldán", 6),
    (54, "Claudia Cortés Lucas", 6),
    (55, "Julián Romero Casado", 6),
    (56, "Patricia Gil Pacheco", 6),

    # Mesa 7: Bonanza (Compañeros de trabajo Novio)
    (57, "David Vargas Luque", 7),
    (58, "Sara Iglesias Nieto", 7),
    (59, "Jorge Gallego Vidal", 7),
    (60, "Cristina Soto Bravo", 7),
    (61, "Emilio Marín Saavedra", 7),
    (62, "Lorena Pascual Mora", 7),
    (63, "Gabriel Redondo León", 7),
    (64, "Nerea Domenech Solís", 7),

    # Mesa 8: San Salvador (Amigos de Sevilla y Cádiz)
    (65, "Juan Carlos Jurado Ríos", 8),
    (66, "Macarena Alarcón Bueno", 8),
    (67, "Enrique Peñalver Cano", 8),
    (68, "Fátima Salgado Bermejo", 8),
    (69, "Samuel Herrero Moya", 8),
    (70, "Irene Carvajal Baeza", 8),
    (71, "Borja Otero Velasco", 8),
    (72, "Victoria Beltrán Rivas", 8),

    # Mesa 9: La Calzada (Amigos de los padres de la novia)
    (73, "Rafael Montoro Carmona", 9),
    (74, "Asunción Barroso Mesa", 9),
    (75, "Joaquín Valiente Luque", 9),
    (76, "Mercedes Carrillo Vera", 9),
    (77, "Pedro Manzano Soria", 9),
    (78, "Concepción Aranda Polo", 9),
    (79, "Miguel Ángel Roldán Gil", 9),
    (80, "Rosario Hurtado Nieto", 9),

    # Mesa 10: Las Piletas (Amigos de los padres del novio)
    (81, "Bernardo Quesada Maroto", 10),
    (82, "Milagros Tejero Pineda", 10),
    (83, "César Palacios Bernabé", 10),
    (84, "Amparo Soler Villalba", 10),
    (85, "Vicente Guisado Román", 10),
    (86, "Josefa Cobos Almansa", 10),
    (87, "Felipe Cordero Zurita", 10),
    (88, "Dolores Perea Maya", 10),

    # Mesa 11: Barrameda (Mesa Joven / Primos pequeños y adolescentes)
    (89, "Ángel Muñoz Prieto", 11),
    (90, "Mireia Madrid Valero", 11),
    (91, "Hugo Gómez Ramos", 11),
    (92, "Vega Barba Moreno", 11),
    (93, "Leo Castro Herrera", 11),
    (94, "Aitana Ortiz Cobo", 11),
    (95, "Darío Quintana Pardo", 11),
    (96, "Noa Santos Guerra", 11),

    # Mesa 12: Tartessos (Invitados especiales y familia lejana de Madrid)
    (97, "Eduardo Blázquez Cuenca", 12),
    (98, "Alicia Mendieta Forner", 12),
    (99, "Cristóbal Cañete Rico", 12),
    (100, "Marina Lucena Tapia", 12),
    (101, "Guillermo Rivas Salcedo", 12),
    (102, "Elisa Villena Osorio", 12),
    (103, "Leandro Sotomayor Alba", 12),
    (104, "Paloma Rebollo Duque", 12)
]

# 1. Generar Excel invitados.xlsx
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Invitados"

# Estilos elegantes para el Excel
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="3A342E", end_color="3A342E", fill_type="solid")
border_thin = Border(
    left=Side(style='thin', color='E0D7CD'),
    right=Side(style='thin', color='E0D7CD'),
    top=Side(style='thin', color='E0D7CD'),
    bottom=Side(style='thin', color='E0D7CD')
)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

headers = ["Número", "Nombre y Apellido", "Número de mesa"]
ws.append(headers)

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_center

for row_idx, guest in enumerate(invitados_data, start=2):
    ws.append(list(guest))
    c1 = ws.cell(row=row_idx, column=1)
    c2 = ws.cell(row=row_idx, column=2)
    c3 = ws.cell(row=row_idx, column=3)
    
    c1.alignment = align_center
    c2.alignment = align_left
    c3.alignment = align_center
    
    c1.border = border_thin
    c2.border = border_thin
    c3.border = border_thin
    
    # Alternar color suave de filas
    if row_idx % 2 == 0:
        row_fill = PatternFill(start_color="FAF7F2", end_color="FAF7F2", fill_type="solid")
        c1.fill = row_fill
        c2.fill = row_fill
        c3.fill = row_fill

ws.column_dimensions['A'].width = 14
ws.column_dimensions['B'].width = 36
ws.column_dimensions['C'].width = 18

wb.save("invitados.xlsx")
print("[OK] Archivo invitados.xlsx generado con exito.")

# 2. Generar JSON invitados.json correspondiente
json_list = []
for guest in invitados_data:
    json_list.append({
        "id": guest[0],
        "nombre": guest[1],
        "mesa": guest[2]
    })

with open("invitados.json", "w", encoding="utf-8") as f:
    json.dump(json_list, f, ensure_ascii=False, indent=2)

print("[OK] Archivo invitados.json generado con exito (104 comensales en 13 mesas).")
